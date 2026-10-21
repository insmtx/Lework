#!/usr/bin/env python3
"""Build the five-sheet payroll workbook from verified, read-only inputs.

This is a presentation exporter: it does not call Payroll or calculate wages.
"""

from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import os
import re
import sys
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import workflow_gate
import person_matches
import attendance_store
import project_binding


COMPONENTS = (
    ("BASE_SALARY", "基础工资"),
    ("PERFORMANCE_SALARY", "绩效工资"),
    ("SENIORITY_SALARY", "工龄工资"),
    ("PROFESSIONAL_TITLE_SALARY", "职称工资"),
    ("CONSTRUCTION_ALLOWANCE", "施工补贴"),
    ("PHONE_ALLOWANCE", "话费补贴"),
    ("COOLING_ALLOWANCE", "降温费"),
    ("TRANSPORT_ALLOWANCE", "交通补贴"),
    ("WEEKEND_OVERTIME_PAY", "双休日加班工资"),
)
COMPONENT_NAMES = dict(COMPONENTS)
DAY_STATUS = {1: "出勤", 2: "事假", 3: "病假", 4: "旷工", 5: "休息", 6: "调整休息", 7: "其他"}
RECONCILIATION = {1: "一致", 2: "人工调整"}
SHEETS = {
    "工资核算明细": ["姓名", "项目", "公司/部门", "岗位", "人员类别", *[name for _, name in COMPONENTS],
                  "历史工资来源月份", "历史应发工资", "核算应发工资", "与历史差额", "状态"],
    "工资基准": ["姓名", "项目", "公司/部门", "岗位", "人员类别", "历史工资来源月份", "历史应发工资",
                 *[value for _, name in COMPONENTS for value in (name, f"{name}来源状态")]],
    "考勤汇总": ["姓名", "考勤记录ID", "项目", "公司/部门", "岗位", "人员类别", "来源出勤天数", "派生出勤天数",
                "确认出勤天数", "事假天数", "病假天数", "旷工天数", "对账状态", "备注"],
    "历史工资对比": ["姓名", "项目", "公司/部门", "岗位", "历史工资来源月份", "历史应发工资", "本次核算应发", "与历史差额"],
    "人工复核事项": ["姓名", "处理状态", "问题类型", "问题说明"],
}
MONEY_FMT = "#,##0.00;(#,##0.00);-"
DAYS_FMT = "0.00;[Red](0.00);-"
WIDTHS = {
    "工资核算明细": [10, 10, 11, 10, 10, *([18] * 13), 10],
    "工资基准": [10, 10, 11, 10, 10, 18, 18, *([18] * 13), 16, 18, 18, 18, 24],
    "考勤汇总": [14, 24, 14, 14, 14, 12, *([14] * 6), 12, 36],
    "历史工资对比": [10, 10, 11, 10, *([18] * 4)],
    "人工复核事项": [12, 10, 18, 36],
}


class ExportError(ValueError):
    """Input evidence is incomplete or inconsistent."""


def _obj(value: object, label: str) -> dict:
    if not isinstance(value, dict):
        raise ExportError(f"{label}_INVALID")
    return value


def _read(path: Path, label: str) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ExportError(f"{label}_UNREADABLE") from exc
    return _obj(value, f"{label}_INVALID")


def _month(value: object) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
        raise ExportError("MONTH_INVALID")
    return value


def _int(value: object, label: str, minimum: int = 0) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < minimum:
        raise ExportError(f"{label}_INVALID")
    return value


def _decimal(value: object, label: str, allow_none: bool = False) -> Decimal | None:
    if value is None and allow_none:
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ExportError(f"{label}_INVALID")
    try:
        amount = Decimal(str(value))
    except InvalidOperation as exc:
        raise ExportError(f"{label}_INVALID") from exc
    if not amount.is_finite():
        raise ExportError(f"{label}_INVALID")
    return amount


def _display_decimal(value: Decimal | None):
    if value is None:
        return None
    return float(value)


def _attrs(record: dict) -> tuple[str, str, str, str, str]:
    name = record.get("name_snapshot")
    if not isinstance(name, str) or not name.strip():
        raise ExportError("ATTENDANCE_NAME_INVALID")
    fields = (record.get("project_snapshot"), record.get("company_snapshot"),
              record.get("position_snapshot"), record.get("employee_type_snapshot"))
    if any(value is not None and not isinstance(value, str) for value in fields[:3]):
        raise ExportError("ATTENDANCE_IDENTITY_FIELD_INVALID")
    if fields[3] is not None and (type(fields[3]) is not int or not 1 <= fields[3] <= 255):
        raise ExportError("ATTENDANCE_IDENTITY_FIELD_INVALID")
    return (name, *fields)


CALCULATION_KEYS = {
    "payroll_month", "batch_id", "operation_id", "precheck_snapshot_id", "valid", "calculation_stage",
    "population_count", "participant_count", "calculated_count", "blocked_count", "batch_blocked",
    "status", "issues", "items", "summary",
}
CALCULATION_ITEM_KEYS = {
    "batch_id", "payroll_person_id", "roster_id", "name", "employee_type", "company_name", "project_name",
    "position", "attendance_id", "attendance_source", "attendance_batch_operation_id", "his_salary_record_id",
    "historical_source_month", "components", "historical_gross_amount", "calculated_gross_amount", "difference",
}
CALCULATION_COMPONENT_KEYS = {"code", "name", "source_present", "source_amount", "calculated_amount"}
CALCULATION_ISSUE_KEYS = {"scope", "payroll_person_id", "roster_id", "name", "code", "message"}
CALCULATION_SUMMARY_KEYS = {
    "population_count", "calculated_count", "historical_gross_total", "calculated_gross_total", "difference_total",
}
ATTENDANCE_RECORD_KEYS = {
    "payroll_person_id", "roster_id", "name_snapshot", "employee_type_snapshot", "company_snapshot",
    "project_snapshot", "source_project_snapshot", "position_snapshot", "source_actual_work_days",
    "derived_actual_work_days", "confirmed_actual_work_days", "personal_leave_days", "sick_leave_days",
    "absent_days", "reconciliation_status", "remark", "items",
}
ATTENDANCE_ITEM_KEYS = {"day", "status", "value", "raw_value"}
SUBMISSION_ROOT_KEYS = {"attendance_month", "submission_mode", "batch_id", "operation_id", "records"}
SUBMISSION_RECORD_KEYS = ATTENDANCE_RECORD_KEYS - {"record_id"}
ATTENDANCE_RECORD_REQUIRED_KEYS = {
    "name_snapshot", "source_actual_work_days", "derived_actual_work_days", "confirmed_actual_work_days",
    "personal_leave_days", "sick_leave_days", "absent_days", "reconciliation_status", "remark", "items",
}


def _keys(value: dict, allowed: set[str], required: set[str], label: str) -> None:
    if set(value) - allowed or required - set(value):
        raise ExportError(f"{label}_STRUCTURE_INVALID")


def _validate_calculation_dto(value: object) -> dict:
    result = _obj(value, "RAW_CALCULATION_RESPONSE")
    if "issues" in result and result["issues"] is None:
        # Go encodes a nil slice without omitempty as JSON null. For a
        # successful response this is equivalent to an empty issue slice.
        result = {**result, "issues": []}
    _keys(result, CALCULATION_KEYS, CALCULATION_KEYS, "RAW_CALCULATION_RESPONSE")
    for key in ("payroll_month", "batch_id", "operation_id", "precheck_snapshot_id", "calculation_stage", "status"):
        if not isinstance(result.get(key), str):
            raise ExportError("RAW_CALCULATION_RESPONSE_FIELD_INVALID")
    for key in ("valid", "batch_blocked"):
        if type(result.get(key)) is not bool:
            raise ExportError("RAW_CALCULATION_RESPONSE_FIELD_INVALID")
    for key in ("population_count", "participant_count", "calculated_count", "blocked_count"):
        _int(result.get(key), f"RAW_{key.upper()}")
    if result.get("status") not in {"SUCCESS", "PARTIAL", "BLOCKED"}:
        raise ExportError("RAW_CALCULATION_STATUS_INVALID")
    items = result.get("items")
    if not isinstance(items, list):
        raise ExportError("RAW_CALCULATION_ITEMS_INVALID")
    for raw in items:
        item = _obj(raw, "RAW_CALCULATION_ITEM")
        _keys(item, CALCULATION_ITEM_KEYS, CALCULATION_ITEM_KEYS - {"batch_id", "attendance_id"}, "RAW_CALCULATION_ITEM")
        for key in ("payroll_person_id", "roster_id", "employee_type", "his_salary_record_id"):
            _int(item.get(key), f"RAW_CALCULATION_ITEM_{key.upper()}", 0)
        for key in ("name", "company_name", "project_name", "position", "historical_source_month",
                    "historical_gross_amount", "calculated_gross_amount", "difference"):
            if not isinstance(item.get(key), str):
                raise ExportError("RAW_CALCULATION_ITEM_FIELD_INVALID")
        for key in ("batch_id", "attendance_source", "attendance_batch_operation_id"):
            if key in item and not isinstance(item[key], str):
                raise ExportError("RAW_CALCULATION_ITEM_FIELD_INVALID")
        if "attendance_id" in item:
            _int(item["attendance_id"], "RAW_CALCULATION_ITEM_ATTENDANCE_ID", 0)
        components = item.get("components")
        if not isinstance(components, list):
            raise ExportError("RAW_CALCULATION_COMPONENTS_INVALID")
        for raw_component in components:
            component = _obj(raw_component, "RAW_CALCULATION_COMPONENT")
            _keys(component, CALCULATION_COMPONENT_KEYS, CALCULATION_COMPONENT_KEYS, "RAW_CALCULATION_COMPONENT")
            if type(component.get("source_present")) is not bool:
                raise ExportError("RAW_CALCULATION_SOURCE_PRESENCE_INVALID")
    issues = result.get("issues")
    if not isinstance(issues, list):
        raise ExportError("RAW_CALCULATION_ISSUES_INVALID")
    for raw_issue in issues:
        issue = _obj(raw_issue, "RAW_CALCULATION_ISSUE")
        _keys(issue, CALCULATION_ISSUE_KEYS, CALCULATION_ISSUE_KEYS - {"payroll_person_id"}, "RAW_CALCULATION_ISSUE")
        for key in ("scope", "name", "code", "message"):
            if not isinstance(issue.get(key), str):
                raise ExportError("RAW_CALCULATION_ISSUE_FIELD_INVALID")
        _int(issue.get("roster_id"), "RAW_CALCULATION_ISSUE_ROSTER_ID", 0)
        if "payroll_person_id" in issue:
            _int(issue["payroll_person_id"], "RAW_CALCULATION_ISSUE_PERSON_ID", 0)
    summary = _obj(result.get("summary"), "RAW_CALCULATION_SUMMARY")
    _keys(summary, CALCULATION_SUMMARY_KEYS, CALCULATION_SUMMARY_KEYS, "RAW_CALCULATION_SUMMARY")
    for key in ("population_count", "calculated_count"):
        _int(summary.get(key), f"RAW_SUMMARY_{key.upper()}")
    for key in ("historical_gross_total", "calculated_gross_total", "difference_total"):
        if not isinstance(summary.get(key), str):
            raise ExportError("RAW_CALCULATION_SUMMARY_FIELD_INVALID")
    return result


def _calculation_dto_from_payroll_response(value: object) -> dict:
    """Normalize a verified Payroll DTO or its actual MCP CallToolResult wrapper."""
    response = _obj(value, "RAW_CALCULATION_RESPONSE")
    if not ({"content", "structuredContent", "isError", "_meta"} & set(response)):
        return _validate_calculation_dto(response)

    allowed = {"content", "structuredContent", "isError", "_meta"}
    if set(response) - allowed or "structuredContent" not in response:
        raise ExportError("RAW_CALCULATION_MCP_ENVELOPE_INVALID")
    if response.get("isError", False) is not False:
        raise ExportError("RAW_CALCULATION_MCP_ERROR")
    if "_meta" in response and not isinstance(response["_meta"], dict):
        raise ExportError("RAW_CALCULATION_MCP_ENVELOPE_INVALID")

    content = response.get("content")
    if not isinstance(content, list) or len(content) != 1:
        raise ExportError("RAW_CALCULATION_MCP_CONTENT_INVALID")
    text_content = _obj(content[0], "RAW_CALCULATION_MCP_CONTENT")
    if set(text_content) != {"type", "text"} or text_content.get("type") != "text" or not isinstance(
        text_content.get("text"), str
    ):
        raise ExportError("RAW_CALCULATION_MCP_CONTENT_INVALID")
    try:
        text_dto = json.loads(text_content["text"])
    except json.JSONDecodeError as exc:
        raise ExportError("RAW_CALCULATION_MCP_CONTENT_INVALID") from exc
    structured_dto = _obj(response["structuredContent"], "RAW_CALCULATION_MCP_STRUCTURED_CONTENT")
    if _canonical_json(text_dto) != _canonical_json(structured_dto):
        raise ExportError("RAW_CALCULATION_MCP_CONTENT_MISMATCH")
    return _validate_calculation_dto(structured_dto)


def _canonical_json(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _decimal_equal(left: object, right: object, label: str) -> bool:
    if left is None or right is None:
        return left is None and right is None
    return _decimal(left, label) == _decimal(right, label)


def _effective_attendance_status(day: dict) -> str:
    confirmed = day.get("user_confirmed_status")
    return confirmed if confirmed is not None else day.get("status")


ATTENDANCE_STATUS_TO_PAYROLL = {
    "work": 1,
    "personal_leave": 2,
    "sick_leave": 3,
    "absent": 4,
    "rest": 5,
    "comp_leave": 6,
}
UNMAPPED_SOURCE_STATUS_LABELS = {
    "transfer": "调离",
    "departed": "离职后",
    "unknown": "未知",
    "blank": "空白",
    "other": "其他",
}


def _attendance_record_projection(person: dict, match: dict, month: str, project_map: dict[str, str],
                                  allow_unmapped_statuses: bool = False) -> dict:
    """Project only facts with a documented one-to-one Payroll DTO representation."""
    name = person.get("name")
    if not isinstance(name, str) or not name.strip():
        raise ExportError("ATTENDANCE_SOURCE_NAME_INVALID")
    raw_project = person.get("project")
    if not isinstance(raw_project, str):
        raise ExportError("ATTENDANCE_SOURCE_PROJECT_INVALID")
    project = project_map.get(raw_project.strip())
    if not project:
        raise ExportError("PROJECT_BINDING_SOURCE_PROJECT_UNBOUND")
    pid = match.get("payroll_person_id") if match.get("status") == "matched" else None
    if pid is not None:
        pid = _person_id(pid, "MATCH_PAYROLL_PERSON_ID")
    roster_id = match.get("roster_id") if match.get("status") == "matched" else None
    if roster_id is not None:
        roster_id = _person_id(roster_id, "MATCH_ROSTER_ID")
    employee_type = person.get("employee_type")
    if employee_type is not None:
        employee_type = _int(employee_type, "ATTENDANCE_EMPLOYEE_TYPE", 1)
        if employee_type > 255:
            raise ExportError("ATTENDANCE_EMPLOYEE_TYPE_INVALID")
    for field in ("company_name", "position"):
        if person.get(field) is not None and not isinstance(person[field], str):
            raise ExportError("ATTENDANCE_SOURCE_IDENTITY_FIELD_INVALID")

    days_in_month = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    daily = person.get("daily")
    if not isinstance(daily, list):
        raise ExportError("ATTENDANCE_SOURCE_DAILY_INVALID")
    by_day = {}
    unmapped_days = {}
    corrections = "confirmed_work_days" in person
    personal = sick = absent = observed = 0
    for day in daily:
        if not isinstance(day, dict) or type(day.get("day")) is not int or not 1 <= day["day"] <= days_in_month:
            raise ExportError("ATTENDANCE_SOURCE_DAY_INVALID")
        if day["day"] in by_day:
            raise ExportError("ATTENDANCE_SOURCE_DAY_DUPLICATE")
        status = _effective_attendance_status(day)
        if not isinstance(status, str) or status not in attendance_store.STATUS:
            raise ExportError("ATTENDANCE_SOURCE_STATUS_INVALID")
        if status not in ATTENDANCE_STATUS_TO_PAYROLL and not allow_unmapped_statuses:
            raise ExportError("ATTENDANCE_STATUS_HAS_NO_PAYROLL_EQUIVALENT")
        raw_mark = day.get("raw_mark")
        if not isinstance(raw_mark, str):
            raise ExportError("ATTENDANCE_SOURCE_RAW_MARK_INVALID")
        confirmed_mark = day.get("user_confirmed_mark")
        if confirmed_mark is not None and confirmed_mark != raw_mark:
            raise ExportError("ATTENDANCE_CONFIRMED_MARK_NOT_REPRESENTABLE")
        if "user_confirmed_status" in day:
            corrections = True
        value = day.get("value")
        if value is not None:
            _decimal(value, "ATTENDANCE_SOURCE_DAILY_VALUE")
        projected = {
            "day": day["day"],
            "status": ATTENDANCE_STATUS_TO_PAYROLL.get(status),
            "value": None if value is None else str(value),
            "raw_value": raw_mark,
        }
        if status not in ATTENDANCE_STATUS_TO_PAYROLL:
            unmapped_days[day["day"]] = status
        by_day[day["day"]] = projected
        observed += int(status == "work")
        personal += int(status == "personal_leave")
        sick += int(status == "sick_leave")
        absent += int(status == "absent")
    if set(by_day) != set(range(1, days_in_month + 1)):
        raise ExportError("ATTENDANCE_SOURCE_DAILY_INCOMPLETE")

    printed = person.get("printed_work_days")
    printed_reliable = (printed not in (None, "") and person.get("row_binding_status") == "VERIFIED"
                        and type(person.get("row_index")) is int
                        and person.get("printed_work_days_row_index") == person.get("row_index"))
    if printed not in (None, "") and not printed_reliable:
        raise ExportError("ATTENDANCE_PRINTED_DAYS_UNBOUND")
    source_days = None if not printed_reliable else _decimal(printed, "ATTENDANCE_PRINTED_DAYS")
    derived_days = Decimal(observed)
    confirmed_days = person.get("confirmed_work_days")
    if confirmed_days is not None:
        confirmed_days = _decimal(confirmed_days, "ATTENDANCE_CONFIRMED_DAYS")
    elif source_days is not None:
        confirmed_days = source_days
    else:
        raise ExportError("ATTENDANCE_CONFIRMED_DAYS_UNAVAILABLE")
    note = person.get("note")
    if note is not None and not isinstance(note, str):
        raise ExportError("ATTENDANCE_NOTE_INVALID")
    return {
        "payroll_person_id": pid,
        "roster_id": roster_id,
        "name_snapshot": name,
        "employee_type_snapshot": employee_type,
        "company_snapshot": person.get("company_name"),
        "project_snapshot": project,
        "source_project_snapshot": raw_project,
        "position_snapshot": person.get("position"),
        "source_actual_work_days": None if source_days is None else str(source_days),
        "derived_actual_work_days": str(derived_days),
        "confirmed_actual_work_days": str(confirmed_days),
        "personal_leave_days": str(personal),
        "sick_leave_days": str(sick),
        "absent_days": str(absent),
        "reconciliation_status": 2 if corrections else 1,
        "remark": note,
        "items": [by_day[day] for day in sorted(by_day)],
        "unmapped_days": unmapped_days,
    }


def _compare_attendance_record(expected: dict, actual: dict) -> None:
    _keys(actual, ATTENDANCE_RECORD_KEYS | {"record_id", "attendance_id", "attendance_source", "attendance_batch_operation_id"},
          ATTENDANCE_RECORD_REQUIRED_KEYS, "ATTENDANCE_SNAPSHOT_RECORD")
    for key in ATTENDANCE_RECORD_KEYS - {"items", "source_actual_work_days", "derived_actual_work_days",
                                         "confirmed_actual_work_days", "personal_leave_days", "sick_leave_days", "absent_days"}:
        if actual.get(key) != expected.get(key):
            raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
    for key in ("source_actual_work_days", "derived_actual_work_days", "confirmed_actual_work_days",
                "personal_leave_days", "sick_leave_days", "absent_days"):
        if not _decimal_equal(actual.get(key), expected.get(key), f"ATTENDANCE_SNAPSHOT_{key.upper()}"):
            raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
    actual_items = actual.get("items")
    if not isinstance(actual_items, list):
        raise ExportError("ATTENDANCE_SNAPSHOT_ITEMS_INVALID")
    by_day = {}
    for raw in actual_items:
        item = _obj(raw, "ATTENDANCE_SNAPSHOT_DAY")
        _keys(item, ATTENDANCE_ITEM_KEYS, ATTENDANCE_ITEM_KEYS, "ATTENDANCE_SNAPSHOT_DAY")
        day = _int(item.get("day"), "ATTENDANCE_SNAPSHOT_DAY", 1)
        if day in by_day:
            raise ExportError("ATTENDANCE_SNAPSHOT_DAY_DUPLICATE")
        by_day[day] = item
    exp_by_day = {item["day"]: item for item in expected["items"]}
    unmapped_days = expected.get("unmapped_days", {})
    if set(by_day) != set(exp_by_day):
        raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
    for day, exp in exp_by_day.items():
        got = by_day[day]
        if day in unmapped_days:
            # Local workbook snapshots use null as an explicit non-Payroll
            # placeholder when an out-of-scope source fact has no DTO mapping.
            # No Payroll status can be inferred from this value.
            if got.get("status") is not None or exp.get("status") is not None:
                raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
        else:
            got_status = _int(got.get("status"), "ATTENDANCE_SNAPSHOT_STATUS", 1)
            if got_status not in DAY_STATUS:
                raise ExportError("ATTENDANCE_STATUS_UNKNOWN")
            if got_status != exp["status"]:
                raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
        if got.get("raw_value") != exp["raw_value"]:
            raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")
        if not _decimal_equal(got.get("value"), exp["value"], "ATTENDANCE_SNAPSHOT_VALUE"):
            raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH")


def _validate_source_bindings(data: dict, gate_path: Path, calculation_path: Path,
                              matches_path: Path, submission_path: Path) -> None:
    raw_calculation = _calculation_dto_from_payroll_response(_read(calculation_path, "RAW_CALCULATION_RESPONSE"))
    calculation = _validate_calculation_dto(_obj(data.get("calculation_result"), "CALCULATION_RESULT"))
    if _canonical_json(raw_calculation) != _canonical_json(calculation):
        raise ExportError("CALCULATION_RESULT_SOURCE_MISMATCH")
    calc_ids = {_person_id(item.get("payroll_person_id"), "CALCULATION_PERSON_ID")
                for item in calculation.get("items", [])}

    gate_request = _read(gate_path, "GATE_REQUEST")
    evidence = _obj(gate_request.get("evidence"), "GATE_EVIDENCE")
    attendance_path = evidence.get("attendance_path")
    if not isinstance(attendance_path, str) or not attendance_path:
        raise ExportError("CURRENT_ATTENDANCE_PATH_REQUIRED")
    try:
        attendance_bytes = Path(attendance_path).read_bytes()
        attendance = json.loads(attendance_bytes)
        matches = _read(matches_path, "PERSON_MATCHES")
        submission = _read(submission_path, "SUBMISSION_REQUEST")
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ExportError("SOURCE_EVIDENCE_UNREADABLE") from exc
    attendance = _obj(attendance, "ATTENDANCE_SOURCE")
    matches = _obj(matches, "PERSON_MATCHES")
    submission = _obj(submission, "SUBMISSION_REQUEST")
    if attendance.get("schema_version") != 1 or not isinstance(attendance.get("people"), list):
        raise ExportError("ATTENDANCE_SOURCE_INVALID")
    month = _month(data["attendance_snapshot"].get("attendance_month"))
    if attendance.get("month") != month:
        raise ExportError("ATTENDANCE_SOURCE_MONTH_MISMATCH")
    attendance_sha = hashlib.sha256(attendance_bytes).hexdigest()
    review = _obj(evidence.get("attendance_review"), "ATTENDANCE_REVIEW")
    if review.get("attendance_sha256") != attendance_sha or review.get("month") != month:
        raise ExportError("ATTENDANCE_REVIEW_SOURCE_MISMATCH")
    roster_path, binding_path = project_binding.default_paths(attendance_path)
    try:
        project_binding_result = project_binding.validate(attendance_path, roster_path, binding_path)
    except project_binding.ProjectBindingError as exc:
        raise ExportError(str(exc)) from exc
    project_map = project_binding_result["project_map"]
    binding_sha = project_binding_result["binding"]["binding_sha256"]
    if (matches.get("schema_version") != 1 or matches.get("matching_rules_version") != person_matches.MATCHING_RULES_VERSION
            or matches.get("attendance_sha256") != attendance_sha
            or matches.get("project_binding_sha256") != binding_sha or not isinstance(matches.get("matches"), list)):
        raise ExportError("PERSON_MATCHES_SOURCE_MISMATCH")
    matches_by_id = {}
    used_ids = set()
    for raw in matches["matches"]:
        item = _obj(raw, "PERSON_MATCH")
        record_id = item.get("record_id")
        if not isinstance(record_id, str) or not record_id or record_id in matches_by_id:
            raise ExportError("PERSON_MATCHES_RECORD_ID_INVALID")
        if item.get("status") not in {"matched", "ambiguous", "unmatched", "pending_confirmation"}:
            raise ExportError("PERSON_MATCH_STATUS_INVALID")
        pid = item.get("payroll_person_id")
        if item.get("status") == "matched":
            pid = _person_id(pid, "MATCH_PAYROLL_PERSON_ID")
            if pid in used_ids:
                raise ExportError("PERSON_MATCHES_PAYROLL_ID_DUPLICATE")
            used_ids.add(pid)
        elif pid is not None:
            raise ExportError("PERSON_MATCH_STATUS_ID_CONFLICT")
        matches_by_id[record_id] = item
    people = attendance["people"]
    if len(matches_by_id) != len(people) or {p.get("record_id") for p in people} != set(matches_by_id):
        raise ExportError("PERSON_MATCHES_ATTENDANCE_SCOPE_MISMATCH")

    snapshot = _obj(data.get("attendance_snapshot"), "ATTENDANCE_SNAPSHOT")
    records = snapshot.get("records")
    if not isinstance(records, list):
        raise ExportError("ATTENDANCE_RECORDS_INVALID")
    snapshot_by_record = {}
    for raw in records:
        record = _obj(raw, "ATTENDANCE_RECORD")
        record_id = record.get("record_id")
        if not isinstance(record_id, str) or not record_id or record_id in snapshot_by_record:
            raise ExportError("ATTENDANCE_SNAPSHOT_RECORD_ID_INVALID")
        snapshot_by_record[record_id] = record
    if set(snapshot_by_record) != set(matches_by_id):
        raise ExportError("ATTENDANCE_SNAPSHOT_SOURCE_SCOPE_MISMATCH")
    source_people_by_record = {}
    for person in people:
        record_id = person.get("record_id")
        match = matches_by_id[record_id]
        source_pid = match.get("payroll_person_id") if match.get("status") == "matched" else None
        calculated = source_pid in calc_ids
        projected = _attendance_record_projection(
            person, match, month, project_map, allow_unmapped_statuses=not calculated,
        )
        actual = snapshot_by_record[record_id]
        if actual.get("payroll_person_id") != projected["payroll_person_id"]:
            raise ExportError("ATTENDANCE_IDENTITY_SOURCE_MISMATCH")
        _compare_attendance_record(projected, actual)
        source_people_by_record[record_id] = person

    _keys(submission, SUBMISSION_ROOT_KEYS, SUBMISSION_ROOT_KEYS, "SUBMISSION_REQUEST")
    if (submission.get("attendance_month") != month or submission.get("submission_mode") != "BATCH_SNAPSHOT"
            or submission.get("batch_id") != calculation.get("batch_id")
            or submission.get("operation_id") != "submit:" + calculation.get("batch_id", "")):
        raise ExportError("SUBMISSION_REQUEST_SCOPE_MISMATCH")
    submit_records = submission.get("records")
    if not isinstance(submit_records, list) or not submit_records:
        raise ExportError("SUBMISSION_REQUEST_RECORDS_INVALID")
    submitted_by_id = {}
    for raw in submit_records:
        record = _obj(raw, "SUBMISSION_REQUEST_RECORD")
        _keys(record, SUBMISSION_RECORD_KEYS, ATTENDANCE_RECORD_REQUIRED_KEYS, "SUBMISSION_REQUEST_RECORD")
        pid = _person_id(record.get("payroll_person_id"), "SUBMISSION_REQUEST_PERSON_ID")
        if pid in submitted_by_id:
            raise ExportError("SUBMISSION_REQUEST_PERSON_DUPLICATE")
        submitted_by_id[pid] = record
    if set(submitted_by_id) != calc_ids:
        raise ExportError("SUBMISSION_REQUEST_SCOPE_MISMATCH")
    for record in records:
        pid = record.get("payroll_person_id")
        if pid is not None and pid in submitted_by_id:
            request_record = submitted_by_id[pid]
            normalized_request = {key: request_record.get(key) for key in SUBMISSION_RECORD_KEYS}
            try:
                _compare_attendance_record(normalized_request, record)
            except ExportError as exc:
                raise ExportError("SUBMISSION_REQUEST_SNAPSHOT_MISMATCH") from exc
    identity = _obj(evidence.get("identity_matching"), "IDENTITY_MATCHING")
    identity_people = identity.get("people")
    if not isinstance(identity_people, list):
        raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
    source_by_id = {person["record_id"]: person for person in people}
    identity_record_ids = identity.get("record_ids")
    if (identity.get("month") != month or identity.get("version") != review.get("version")
            or identity.get("attendance_sha256") != attendance_sha
            or not isinstance(identity_record_ids, list)
            or any(not isinstance(record_id, str) for record_id in identity_record_ids)
            or len(identity_record_ids) != len(set(identity_record_ids))
            or set(identity_record_ids) != set(source_by_id)):
        raise ExportError("IDENTITY_MAPPING_SOURCE_SCOPE_MISMATCH")
    expected_identity = {
        record_id: (match["status"], match.get("payroll_person_id"), source_by_id[record_id].get("name"))
        for record_id, match in matches_by_id.items()
    }
    actual_identity = {}
    for item in identity_people:
        item = _obj(item, "IDENTITY_MAPPING_PERSON")
        record_id, name, status = item.get("record_id"), item.get("name"), item.get("status")
        pid = item.get("payroll_person_id")
        if (not isinstance(record_id, str) or record_id not in source_by_id or record_id in actual_identity
                or name != source_by_id[record_id].get("name")
                or status not in {"matched", "unmatched", "ambiguous", "pending_confirmation"}):
            raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
        if status == "matched":
            pid = _person_id(pid, "IDENTITY_MAPPING_PERSON_ID")
        elif pid is not None:
            raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
        actual_identity[record_id] = (status, pid, name)
    if set(actual_identity) != set(expected_identity):
        raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
    if actual_identity != expected_identity:
        raise ExportError("IDENTITY_MAPPING_SOURCE_MISMATCH")
    matched_by_pid = {
        pid: (record_id, name)
        for record_id, (status, pid, name) in actual_identity.items()
        if status == "matched"
    }
    if len(matched_by_pid) != sum(status == "matched" for status, _, _ in actual_identity.values()):
        raise ExportError("IDENTITY_MAPPING_PERSON_ID_DUPLICATE")
    selected_identity = {
        (matched_by_pid[pid][0], pid, matched_by_pid[pid][1])
        for pid in calc_ids if pid in matched_by_pid
    }
    expected_calculation_identity = {
        (record_id, record.get("payroll_person_id"), record.get("name_snapshot"))
        for record_id, record in snapshot_by_record.items()
        if record.get("payroll_person_id") in calc_ids
    }
    if selected_identity != expected_calculation_identity or len(selected_identity) != len(calc_ids):
        raise ExportError("IDENTITY_CALCULATION_SCOPE_MISMATCH")
    submission_receipt = _obj(evidence.get("submission_receipt"), "SUBMISSION_RECEIPT")
    payload_sha = submission_receipt.get("payload_sha256")
    if not isinstance(payload_sha, str) or not re.fullmatch(r"[0-9a-f]{64}", payload_sha):
        raise ExportError("SUBMISSION_PAYLOAD_SHA256_INVALID")
    receipt_ids = submission_receipt.get("payroll_person_ids")
    if not isinstance(receipt_ids, list) or set(receipt_ids) != set(submitted_by_id) or len(receipt_ids) != len(submitted_by_id):
        raise ExportError("SUBMISSION_RECEIPT_SCOPE_MISMATCH")
    return source_people_by_record


def _person_id(value: object, label: str) -> int:
    return _int(value, label, 1)


def _calculate_rows(result: dict, month: str, batch_id: str, operation_id: str, snapshot_ids: set[int]) -> dict[int, dict]:
    if result.get("status") != "SUCCESS" or result.get("valid") is not True or result.get("batch_blocked") is not False:
        raise ExportError("CALCULATION_NOT_SUCCESSFUL")
    if result.get("payroll_month") != month or result.get("batch_id") != batch_id or result.get("operation_id") != operation_id:
        raise ExportError("CALCULATION_SCOPE_MISMATCH")
    people = result.get("items")
    if not isinstance(people, list) or not people:
        raise ExportError("CALCULATION_ITEMS_INVALID")
    population = _int(result.get("population_count"), "POPULATION_COUNT")
    participants = _int(result.get("participant_count"), "PARTICIPANT_COUNT")
    calculated = _int(result.get("calculated_count"), "CALCULATED_COUNT")
    blocked = _int(result.get("blocked_count"), "BLOCKED_COUNT")
    if not (population == participants == calculated == len(people)) or blocked != 0:
        raise ExportError("CALCULATION_COUNTS_INCONSISTENT")
    by_id = {}
    for raw in people:
        item = _obj(raw, "CALCULATION_ITEM")
        pid = _person_id(item.get("payroll_person_id"), "CALCULATION_PERSON_ID")
        if pid in by_id or pid not in snapshot_ids:
            raise ExportError("CALCULATION_PERSON_SCOPE_INVALID")
        if item.get("batch_id") != batch_id:
            raise ExportError("CALCULATION_ITEM_BATCH_MISMATCH")
        if item.get("attendance_batch_operation_id") != operation_id.replace("calculate:", "submit:", 1):
            raise ExportError("CALCULATION_ITEM_ATTENDANCE_OPERATION_MISMATCH")
        components = item.get("components")
        if not isinstance(components, list):
            raise ExportError("COMPONENTS_INVALID")
        component_map = {}
        for component_raw in components:
            component = _obj(component_raw, "COMPONENT")
            code = component.get("code")
            if code not in COMPONENT_NAMES or code in component_map or component.get("name") != COMPONENT_NAMES[code]:
                raise ExportError("COMPONENT_CODE_INVALID_OR_DUPLICATE")
            present = component.get("source_present")
            if present is not True and present is not False and present is not None:
                raise ExportError("COMPONENT_SOURCE_PRESENCE_INVALID")
            source_amount = _decimal(component.get("source_amount"), "SOURCE_AMOUNT", allow_none=present is not True)
            if present and source_amount is None:
                raise ExportError("COMPONENT_SOURCE_AMOUNT_MISSING")
            if present is False and source_amount is not None:
                raise ExportError("COMPONENT_SOURCE_AMOUNT_CONTRADICTS_PRESENCE")
            component_map[code] = {
                "source_present": present,
                "source_amount": source_amount,
                "calculated_amount": _decimal(component.get("calculated_amount"), "CALCULATED_AMOUNT"),
            }
        if set(component_map) != set(COMPONENT_NAMES):
            raise ExportError("COMPONENT_SET_INCOMPLETE")
        gross = _decimal(item.get("calculated_gross_amount"), "CALCULATED_GROSS")
        gross_sum = sum((entry["calculated_amount"] for entry in component_map.values()), Decimal("0"))
        if gross != gross_sum:
            raise ExportError("CALCULATION_GROSS_COMPONENT_MISMATCH")
        historical = _decimal(item.get("historical_gross_amount"), "HISTORICAL_GROSS", allow_none=True)
        source_month = item.get("historical_source_month")
        history_id = item.get("his_salary_record_id")
        difference = _decimal(item.get("difference"), "DIFFERENCE", allow_none=True)
        if historical is None:
            if source_month is not None or history_id is not None or difference is not None:
                raise ExportError("HISTORICAL_SALARY_EVIDENCE_INCOMPLETE")
        else:
            source_month = _month(source_month)
            if source_month > month or history_id is None or difference is None or gross - historical != difference:
                raise ExportError("HISTORICAL_SALARY_EVIDENCE_INVALID")
            _int(history_id, "HISTORICAL_SALARY_RECORD_ID", 1)
        by_id[pid] = {**item, "components_by_code": component_map, "calculated_decimal": gross,
                      "historical_decimal": historical, "difference_decimal": difference,
                      "historical_source_month": source_month}
    summary = _obj(result.get("summary"), "CALCULATION_SUMMARY")
    if _int(summary.get("calculated_count"), "SUMMARY_CALCULATED_COUNT") != calculated or _int(
        summary.get("population_count"), "SUMMARY_POPULATION_COUNT"
    ) != population:
        raise ExportError("CALCULATION_SUMMARY_COUNT_MISMATCH")
    total = sum((item["calculated_decimal"] for item in by_id.values()), Decimal("0"))
    history_values = [item["historical_decimal"] for item in by_id.values() if item["historical_decimal"] is not None]
    differences = [item["difference_decimal"] for item in by_id.values() if item["difference_decimal"] is not None]
    for key, expected in (("calculated_gross_total", total), ("historical_gross_total", sum(history_values, Decimal("0"))),
                          ("difference_total", sum(differences, Decimal("0")))):
        if _decimal(summary.get(key), f"SUMMARY_{key.upper()}") != expected:
            raise ExportError("CALCULATION_SUMMARY_AMOUNT_MISMATCH")
    return by_id


def _validate_gate(path: Path, month: str, calculation: dict) -> dict:
    try:
        request = _read(path, "GATE_REQUEST")
    except ExportError:
        raise
    if request.get("action") != "PAYROLL_EXPORT":
        raise ExportError("PAYROLL_EXPORT_GATE_REQUIRED")
    gate = workflow_gate.evaluate(request)
    if gate.get("ok") is not True or gate.get("next_action") != "GENERATE_AND_DECLARE_REAL_PAYROLL_OUTPUT":
        raise ExportError("PAYROLL_EXPORT_GATE_BLOCKED")
    evidence = request.get("evidence", {})
    receipt = _obj(evidence.get("calculation_receipt"), "CALCULATION_RECEIPT")
    submission = _obj(evidence.get("submission_receipt"), "SUBMISSION_RECEIPT")
    precheck = _obj(evidence.get("payroll_precheck"), "PAYROLL_PRECHECK")
    scope_review = _obj(evidence.get("submission_scope_review"), "SUBMISSION_SCOPE_REVIEW")
    batch = calculation.get("batch_id")
    operation = calculation.get("operation_id")
    scope = evidence.get("calculation_receipt")
    if not isinstance(scope, dict) or scope.get("month") != month:
        raise ExportError("GATE_CALCULATION_MONTH_MISMATCH")
    if receipt.get("status") != "SUCCEEDED" or receipt.get("batch_id") != batch or receipt.get("operation_id") != operation:
        raise ExportError("GATE_CALCULATION_RECEIPT_MISMATCH")
    if submission.get("status") != "SUCCEEDED" or submission.get("batch_id") != batch:
        raise ExportError("GATE_SUBMISSION_RECEIPT_MISMATCH")
    if submission.get("operation_id") != calculation.get("items", [{}])[0].get("attendance_batch_operation_id"):
        raise ExportError("GATE_SUBMISSION_OPERATION_MISMATCH")
    calc_by_id = {item["payroll_person_id"]: item["name"] for item in calculation["items"]}
    if len(calc_by_id) != len(calculation["items"]):
        raise ExportError("CALCULATION_PERSON_SCOPE_INVALID")
    expected_people = None
    for evidence_part in (receipt, precheck, scope_review):
        people = evidence_part.get("people")
        if not isinstance(people, list):
            raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
        mapped = []
        for person in people:
            if not isinstance(person, dict):
                raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
            pid = person.get("payroll_person_id")
            record_id, name = person.get("record_id"), person.get("name")
            if (pid not in calc_by_id or not isinstance(record_id, str) or not record_id.strip()
                    or name != calc_by_id[pid]):
                raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
            mapped.append((record_id, pid, name))
        if len(mapped) != len(set(mapped)) or {pid for _, pid, _ in mapped} != set(calc_by_id):
            raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
        if expected_people is not None and set(mapped) != expected_people:
            raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
        expected_people = set(mapped)

    # Identity evidence covers the whole attendance document. The three
    # execution records above cover only the authorized/calculated subset.
    # Keep those contracts separate, then bind the calculated subset back to
    # unique, positively matched identity rows.
    identity = _obj(evidence.get("identity_matching"), "IDENTITY_MATCHING")
    identity_people = identity.get("people")
    attendance_review = _obj(evidence.get("attendance_review"), "ATTENDANCE_REVIEW")
    review_people = attendance_review.get("people")
    identity_record_ids = identity.get("record_ids")
    review_record_ids = attendance_review.get("record_ids")
    if (not isinstance(identity_people, list) or not isinstance(review_people, list)
            or identity.get("month") != attendance_review.get("month")
            or identity.get("version") != attendance_review.get("version")
            or identity.get("attendance_sha256") != attendance_review.get("attendance_sha256")
            or not isinstance(identity_record_ids, list) or not isinstance(review_record_ids, list)
            or any(not isinstance(record_id, str) or not record_id.strip()
                   for record_id in identity_record_ids + review_record_ids)
            or len(identity_record_ids) != len(set(identity_record_ids))
            or len(review_record_ids) != len(set(review_record_ids))
            or set(identity_record_ids) != set(review_record_ids)):
        raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
    review_by_record = {}
    for person in review_people:
        person = _obj(person, "ATTENDANCE_REVIEW_PERSON")
        record_id, name = person.get("record_id"), person.get("name")
        if not isinstance(record_id, str) or not record_id.strip() or record_id in review_by_record:
            raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
        review_by_record[record_id] = name
    if set(review_by_record) != set(review_record_ids):
        raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")

    identities_by_record = {}
    matched_by_pid = {}
    for raw in identity_people:
        person = _obj(raw, "IDENTITY_MAPPING_PERSON")
        record_id, name, status = person.get("record_id"), person.get("name"), person.get("status")
        pid = person.get("payroll_person_id")
        if (not isinstance(record_id, str) or record_id not in review_by_record
                or record_id in identities_by_record or name != review_by_record[record_id]
                or status not in {"matched", "unmatched", "ambiguous", "pending_confirmation"}):
            raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
        if status == "matched":
            pid = _person_id(pid, "IDENTITY_PAYROLL_PERSON_ID")
            if pid in matched_by_pid:
                raise ExportError("IDENTITY_MAPPING_PERSON_ID_DUPLICATE")
            matched_by_pid[pid] = (record_id, name)
        elif pid is not None:
            raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")
        identities_by_record[record_id] = (status, pid, name)
    if set(identities_by_record) != set(review_by_record):
        raise ExportError("IDENTITY_MAPPING_EVIDENCE_INVALID")

    calculated_people = set()
    for pid, name in calc_by_id.items():
        identity_row = matched_by_pid.get(pid)
        if identity_row is None or identity_row[1] != name:
            raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
        calculated_people.add((identity_row[0], pid, name))
    if calculated_people != expected_people:
        raise ExportError("GATE_CALCULATION_PEOPLE_MISMATCH")
    if precheck.get("batch_id") != batch or precheck.get("snapshot_id", calculation.get("precheck_snapshot_id")) != calculation.get("precheck_snapshot_id"):
        raise ExportError("GATE_PRECHECK_BINDING_MISMATCH")
    return gate


def _review_items(attendance: list[dict], calc_by_id: dict[int, dict], month: str, issues: object) -> list[list[object]]:
    grouped: dict[tuple[str, str], dict] = {}

    def add(identity: object, name: object, status: str, category: str, detail: str) -> None:
        stable = str(identity) if identity not in (None, "") else f"name:{name or ''}"
        key = (stable, category)
        entry = grouped.setdefault(key, {"name": str(name or ""), "status": status, "details": []})
        if detail not in entry["details"]:
            entry["details"].append(detail)

    for record in attendance:
        pid = record.get("payroll_person_id")
        identity = f"payroll:{pid}" if pid is not None else f"record:{record.get('record_id') or ''}"
        name = record.get("name_snapshot") or ""
        if pid is None or pid not in calc_by_id:
            add(identity, name, "待处理", "未纳入核算", "该考勤人员未出现在本次成功核算明细中。")
        if record.get("reconciliation_status") == 2:
            add(identity, name, "待确认", "考勤人工调整", str(record.get("remark") or "来源记录标记为人工调整。"))
        elif record.get("reconciliation_status") not in RECONCILIATION:
            add(identity, name, "待核实", "考勤对账状态未知", "来源考勤快照未提供可识别的对账状态。")
        items = record.get("items", [])
        if not isinstance(items, list):
            continue
        days = [entry.get("day") for entry in items if isinstance(entry, dict)]
        if len(days) != len(set(days)):
            add(identity, name, "待核实", "考勤日期重复", "重复日期：" + ", ".join(map(str, sorted(set(day for day in days if days.count(day) > 1)))))
        expected = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
        if set(days) != set(range(1, expected + 1)):
            missing = sorted(set(range(1, expected + 1)) - set(days))
            add(identity, name, "待核实", "考勤日明细不完整", "缺少日期：" + ", ".join(map(str, missing)))
        observed_work = sum(1 for entry in items if isinstance(entry, dict) and entry.get("status") == 1)
        source_days = record.get("source_actual_work_days")
        derived_days = record.get("derived_actual_work_days")
        if source_days is not None and _decimal(source_days, "SOURCE_WORK_DAYS") != Decimal(observed_work):
            add(identity, name, "待核实", "印刷出勤与逐日观察不一致",
                f"来源印刷值为 {source_days}，逐日出勤标记计数为 {observed_work}；两项原始证据均予保留。")
        if derived_days is not None and _decimal(derived_days, "DERIVED_WORK_DAYS") != Decimal(observed_work):
            add(identity, name, "待核实", "派生出勤与逐日观察不一致",
                f"派生值为 {derived_days}，逐日出勤标记计数为 {observed_work}；未据此覆盖任一值。")
    calc_name_by_id = {pid: item.get("name", "") for pid, item in calc_by_id.items()}
    calc_id_by_roster = {item.get("roster_id"): pid for pid, item in calc_by_id.items()}
    if isinstance(issues, list):
        for issue_index, issue in enumerate(issues):
            if isinstance(issue, dict):
                pid = issue.get("payroll_person_id")
                roster_id = issue.get("roster_id")
                if pid is None and roster_id in calc_id_by_roster:
                    pid = calc_id_by_roster[roster_id]
                if pid is not None:
                    identity = f"payroll:{pid}"
                elif roster_id is not None:
                    identity = f"roster:{roster_id}"
                else:
                    identity = f"issue:{issue.get('code') or '核算事项'}:{issue_index}"
                name = issue.get("name") or calc_name_by_id.get(pid, "")
                category = str(issue.get("code") or "核算事项")
                scope = str(issue.get("scope") or "")
                message = str(issue.get("message") or issue.get("reason") or "Payroll 返回了需处理事项。")
                issue_context = []
                if scope:
                    issue_context.append(f"范围：{scope}")
                if roster_id is not None:
                    issue_context.append(f"花名册ID：{roster_id}")
                detail = "；".join([*issue_context, message])
                add(identity, name, "待处理", category, detail)
            else:
                add("issue:核算事项", "", "待处理", "核算事项", str(issue))
    return [[identity, entry["name"], entry["status"], category, "；".join(entry["details"])]
            for (identity, category), entry in sorted(grouped.items())]


def _build_rows(data: dict, gate_request: Path, source_people_by_record: dict[str, dict] | None = None
                ) -> tuple[str, list[list], list[list], list[list], list[list], list[list]]:
    workflow = _obj(data.get("workflow_state"), "WORKFLOW_STATE")
    precheck = _obj(data.get("latest_precheck"), "LATEST_PRECHECK")
    snapshot = _obj(data.get("attendance_snapshot"), "ATTENDANCE_SNAPSHOT")
    calc = _obj(data.get("calculation_result"), "CALCULATION_RESULT")
    month = _month(snapshot.get("attendance_month"))
    if calc.get("payroll_month") != month or precheck.get("valid") is not True or precheck.get("batch_blocked") is not False:
        raise ExportError("PRECHECK_OR_MONTH_INVALID")
    if workflow.get("phase") != "CALCULATED_ARTIFACT_PENDING" or workflow.get("calculation_completed") is not True:
        raise ExportError("WORKFLOW_NOT_READY_FOR_EXPORT")
    if workflow.get("human_precheck_confirmed") is not True or workflow.get("precheck_confirmed_snapshot_id") != precheck.get("snapshot_id"):
        raise ExportError("PRECHECK_CONFIRMATION_MISMATCH")
    if workflow.get("precheck_snapshot_id") != precheck.get("snapshot_id"):
        raise ExportError("WORKFLOW_PRECHECK_SNAPSHOT_MISMATCH")
    if _int(workflow.get("attendance_generation"), "ATTENDANCE_GENERATION") != _int(
        workflow.get("attendance_approved_generation"), "ATTENDANCE_APPROVED_GENERATION"
    ):
        raise ExportError("ATTENDANCE_GENERATION_UNAPPROVED")
    attendance = snapshot.get("records")
    if not isinstance(attendance, list) or not attendance:
        raise ExportError("ATTENDANCE_RECORDS_INVALID")
    snapshot_by_id = {}
    for record_raw in attendance:
        record = _obj(record_raw, "ATTENDANCE_RECORD")
        pid = record.get("payroll_person_id")
        if pid is not None:
            pid = _person_id(pid, "ATTENDANCE_PAYROLL_PERSON_ID")
            if pid in snapshot_by_id:
                raise ExportError("ATTENDANCE_PERSON_ID_DUPLICATE")
        _attrs(record)
        if not isinstance(record.get("items"), list):
            raise ExportError("ATTENDANCE_DAILY_ITEMS_INVALID")
        if pid is not None:
            snapshot_by_id[pid] = record
    batch_id = calc.get("batch_id")
    operation_id = calc.get("operation_id")
    if not isinstance(batch_id, str) or not batch_id or not isinstance(operation_id, str):
        raise ExportError("CALCULATION_OPERATION_REQUIRED")
    calc_by_id = _calculate_rows(calc, month, batch_id, operation_id, set(snapshot_by_id))
    if (_int(precheck.get("participant_count"), "PRECHECK_PARTICIPANT_COUNT") != len(calc_by_id)
            or _int(precheck.get("calculated_count"), "PRECHECK_CALCULATED_COUNT") != len(calc_by_id)
            or _int(precheck.get("blocked_count"), "PRECHECK_BLOCKED_COUNT") != 0):
        raise ExportError("PRECHECK_SCOPE_MISMATCH")
    if precheck.get("snapshot_id") != calc.get("precheck_snapshot_id"):
        raise ExportError("CALCULATION_PRECHECK_SNAPSHOT_MISMATCH")
    _validate_gate(gate_request, month, calc)

    detail, baseline, attendance_rows, history, reviews = [], [], [], [], []
    for pid, item in sorted(calc_by_id.items()):
        source = snapshot_by_id[pid]
        name, project, company, position, employee_type = _attrs(source)
        calc_name = item.get("name")
        if not isinstance(calc_name, str) or calc_name != name:
            raise ExportError("ATTENDANCE_CALCULATION_IDENTITY_MISMATCH")
        if any(item.get(field) != expected for field, expected in (("project_name", project), ("company_name", company), ("position", position)) if expected is not None):
            raise ExportError("CALCULATION_IDENTITY_FIELDS_MISMATCH")
        comp = item["components_by_code"]
        amounts = [_display_decimal(comp[code]["calculated_amount"]) for code, _ in COMPONENTS]
        historical = item["historical_decimal"]
        gross = item["calculated_decimal"]
        difference = item["difference_decimal"]
        history_month = item["historical_source_month"]
        detail.append([name, project, company, position, employee_type, *amounts, history_month,
                       _display_decimal(historical), _display_decimal(gross), _display_decimal(difference), "已核算"])
        base_row = [name, project, company, position, employee_type, history_month, _display_decimal(historical)]
        for code, _ in COMPONENTS:
            part = comp[code]
            source_status = {True: "是", False: "否", None: "待核实"}[part["source_present"]]
            base_row.extend([_display_decimal(part["source_amount"]), source_status])
            calculated_amount = part["calculated_amount"]
            source_amount = part["source_amount"]
            if part["source_present"] is False and calculated_amount != 0:
                reviews.append([f"payroll:{pid}", name, "待核实", "工资来源与计算金额不一致",
                                f"{COMPONENT_NAMES[code]}无来源记录但计算金额为 {_display_decimal(calculated_amount)}；以 Payroll 原始结果为准，需核实。"])
            elif part["source_present"] is None:
                reviews.append([f"payroll:{pid}", name, "待核实", "工资分项来源状态未知",
                                f"{COMPONENT_NAMES[code]}来源状态未知，来源金额为 {_display_decimal(source_amount)}，计算金额为 {_display_decimal(calculated_amount)}。"])
        baseline.append(base_row)
        if historical is not None:
            history.append([name, project, company, position, history_month, _display_decimal(historical),
                            _display_decimal(gross), _display_decimal(difference)])
        else:
            reviews.append([f"payroll:{pid}", name, "待核实", "历史工资来源缺失", "Payroll 核算结果未提供历史工资来源记录，未生成历史差额。"])
    # Preserve every attendance fact, including people not included in Payroll calculation.
    for source_raw in sorted(attendance, key=lambda item: str(item.get("record_id", ""))):
        source = _obj(source_raw, "ATTENDANCE_RECORD")
        pid = source.get("payroll_person_id")
        name, project, company, position, employee_type = _attrs(source)
        record_id = source.get("record_id")
        if not isinstance(record_id, str) or not record_id.strip():
            raise ExportError("ATTENDANCE_RECORD_ID_INVALID")
        by_day = {}
        for day_item_raw in source["items"]:
            day_item = _obj(day_item_raw, "ATTENDANCE_DAY")
            day = _int(day_item.get("day"), "ATTENDANCE_DAY", 1)
            if day > calendar.monthrange(int(month[:4]), int(month[5:]))[1]:
                raise ExportError("ATTENDANCE_DAY_OUT_OF_MONTH")
            if day in by_day:
                raise ExportError("ATTENDANCE_DAY_DUPLICATE")
            status = day_item.get("status")
            if status is not None:
                status = _int(status, "ATTENDANCE_STATUS", 1)
                if status not in DAY_STATUS:
                    raise ExportError("ATTENDANCE_STATUS_UNKNOWN")
            by_day[day] = day_item
        days_in_month = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
        if set(by_day) != set(range(1, days_in_month + 1)):
            raise ExportError("ATTENDANCE_SOURCE_DAILY_INCOMPLETE")
        row = [name, record_id, project, company, position, employee_type,
               source.get("source_actual_work_days"), source.get("derived_actual_work_days"),
               source.get("confirmed_actual_work_days"), source.get("personal_leave_days"),
               source.get("sick_leave_days"), source.get("absent_days"),
               RECONCILIATION.get(source.get("reconciliation_status"), "待核实"), source.get("remark")]
        source_person = (source_people_by_record or {}).get(record_id)
        source_day_by_day = {
            source_item["day"]: source_item
            for source_item in source_person.get("daily", [])
            if isinstance(source_item, dict) and type(source_item.get("day")) is int
        } if source_person is not None else {}
        for day in range(1, days_in_month + 1):
            item = by_day[day]
            source_day = source_day_by_day.get(day)
            source_status = _effective_attendance_status(source_day) if source_day is not None else None
            out_of_scope = pid is None or pid not in calc_by_id
            if (out_of_scope and source_status is not None
                    and source_status not in ATTENDANCE_STATUS_TO_PAYROLL):
                label = UNMAPPED_SOURCE_STATUS_LABELS.get(source_status, "未映射")
                value_raw = source_day.get("value")
                raw_mark = source_day.get("raw_mark")
                value = "（空）" if value_raw is None else str(value_raw)
                mark = "（空）" if raw_mark is None else str(raw_mark)
                row.append(f"待核实（来源状态：{label}）\n值：{value}\n原始：{mark}")
                continue
            status = _int(item.get("status"), "ATTENDANCE_STATUS", 1)
            if status not in DAY_STATUS:
                raise ExportError("ATTENDANCE_STATUS_UNKNOWN")
            value = "（空）" if item.get("value") is None else str(item.get("value"))
            raw_mark = "（空）" if item.get("raw_value") is None else str(item.get("raw_value"))
            row.append(f"{DAY_STATUS[status]}\n值：{value}\n原始：{raw_mark}")
        attendance_rows.append(row)
    reviews.extend(_review_items(attendance, calc_by_id, month, calc.get("issues")))
    if source_people_by_record:
        snapshot_by_record = {record["record_id"]: record for record in attendance}
        for record_id, source_person in source_people_by_record.items():
            source_record = snapshot_by_record[record_id]
            pid = source_record.get("payroll_person_id")
            if pid is not None and pid in calc_by_id:
                continue
            name = source_record.get("name_snapshot") or source_person.get("name") or ""
            for item in source_person.get("daily", []):
                if not isinstance(item, dict):
                    continue
                status = _effective_attendance_status(item)
                if status in ATTENDANCE_STATUS_TO_PAYROLL:
                    continue
                label = UNMAPPED_SOURCE_STATUS_LABELS.get(status, "未映射")
                reviews.append([
                    f"record:{record_id}", name, "待核实", "核算范围外考勤状态待核实",
                    f"{int(item['day'])}日来源状态为{label}，未映射为 Payroll 日状态；原始标记：{item.get('raw_mark') or '（空）'}。",
                ])
    # Component-level findings use the same stable identity + issue category aggregation.
    if reviews:
        grouped_reviews = {}
        for identity, name, status, category, detail_text in reviews:
            key = (str(identity), str(category))
            group = grouped_reviews.setdefault(key, {"name": name, "status": status, "details": []})
            if detail_text not in group["details"]:
                group["details"].append(detail_text)
        reviews = []
        for (identity, category), group in sorted(grouped_reviews.items()):
            if identity.startswith("payroll:"):
                trace = f"Payroll Person ID {identity.split(':', 1)[1]}"
            elif identity.startswith("record:"):
                trace = f"考勤记录ID {identity.split(':', 1)[1]}"
            elif identity.startswith("roster:"):
                trace = f"花名册ID {identity.split(':', 1)[1]}"
            else:
                trace = ""
            detail_text = "；".join(group["details"])
            if trace:
                detail_text = f"追溯标识：{trace}；{detail_text}"
            reviews.append([group["name"], group["status"], category, detail_text])
    if not reviews:
        reviews = [[None, "已完成", "无待处理事项", "本次暂无待处理事项"]]
    return month, detail, baseline, attendance_rows, history, reviews


def _sheet_headers(name: str, month: str) -> list[str]:
    if name != "考勤汇总":
        return SHEETS[name]
    days = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    return [*SHEETS[name], *[f"{day}日" for day in range(1, days + 1)]]


def _style_sheet(ws, name: str, subtitle: str, rows: list[list], month: str) -> None:
    headers = _sheet_headers(name, month)
    last_col = len(headers)
    ws.append([name])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    blue = "17365D"
    fill = PatternFill("solid", fgColor="D9EAF7")
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color=blue)
    ws["A1"].fill = fill
    ws["A2"].font = Font(name="Arial", size=10, bold=True, color="404040")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    for cell in ws[3]:
        cell.font = Font(name="Arial", bold=True, color=blue)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "O4" if name == "考勤汇总" else "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(last_col)}{max(3, ws.max_row)}"
    widths = [*WIDTHS[name], *([8] * (last_col - len(WIDTHS[name])))]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if name == "工资核算明细":
        for row in range(4, ws.max_row + 1):
            for col in range(6, 15):
                ws.cell(row, col).number_format = MONEY_FMT
            for col in (16, 17, 18):
                ws.cell(row, col).number_format = MONEY_FMT
    elif name == "工资基准":
        for row in range(4, ws.max_row + 1):
            for col in (7, *range(8, 25, 2)):
                ws.cell(row, col).number_format = MONEY_FMT
    elif name == "历史工资对比":
        for row in range(4, ws.max_row + 1):
            for col in range(6, 9):
                ws.cell(row, col).number_format = MONEY_FMT
    elif name == "考勤汇总":
        for row in range(4, ws.max_row + 1):
            for col in range(7, 13):
                ws.cell(row, col).number_format = DAYS_FMT
            for col in range(15, last_col + 1):
                ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 42
    elif name == "人工复核事项":
        for row in range(4, ws.max_row + 1):
            ws.cell(row, 4).alignment = Alignment(vertical="top", wrap_text=True)


def export(input_path: Path, gate_path: Path, output_path: Path, calculation_response_path: Path,
           person_matches_path: Path, submission_request_path: Path) -> dict:
    source_paths = {input_path.resolve(), gate_path.resolve(), calculation_response_path.resolve(),
                    person_matches_path.resolve(), submission_request_path.resolve()}
    if output_path.suffix.lower() != ".xlsx" or output_path.exists() or output_path.is_symlink() or output_path.resolve() in source_paths:
        raise ExportError("OUTPUT_PATH_INVALID_OR_EXISTS")
    if not output_path.parent.is_dir():
        raise ExportError("OUTPUT_PARENT_MISSING")
    data = _read(input_path, "INPUT")
    source_people_by_record = _validate_source_bindings(
        data, gate_path, calculation_response_path, person_matches_path, submission_request_path,
    )
    month, detail, baseline, attendance, history, reviews = _build_rows(data, gate_path, source_people_by_record)
    wb = Workbook()
    wb.remove(wb.active)
    row_sets = (detail, baseline, attendance, history, reviews)
    for (name, _), rows in zip(SHEETS.items(), row_sets):
        ws = wb.create_sheet(name)
        count = len(detail)
        _style_sheet(ws, name, f"{month}工资核算成功，共核算 {count} 人。", rows, month)
    fd, temp_name = tempfile.mkstemp(prefix=f".{output_path.stem}-", suffix=".xlsx", dir=output_path.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        wb.save(temp_path)
        check = load_workbook(temp_path, read_only=False, data_only=False)
        if check.sheetnames != list(SHEETS) or any(check[name].max_column != len(_sheet_headers(name, month))
                                                   for name in SHEETS):
            raise ExportError("OUTPUT_VALIDATION_FAILED")
        if any(cell.data_type == "f" for ws in check.worksheets for row in ws.iter_rows() for cell in row):
            raise ExportError("UNEXPECTED_FORMULA")
        check.close()
        try:
            os.link(temp_path, output_path)
        except FileExistsError as exc:
            raise ExportError("OUTPUT_PATH_INVALID_OR_EXISTS") from exc
        temp_path.unlink()
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return {"ok": True, "status": "EXPORTED", "path": str(output_path.resolve()), "month": month,
            "calculated_people": len(detail), "attendance_people_rows": len(attendance),
            "attendance_day_rows": sum(len(record.get("items", [])) for record in data["attendance_snapshot"]["records"]),
            "review_items": 0 if reviews == [[None, "已完成", "无待处理事项", "本次暂无待处理事项"]] else len(reviews)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="verified payroll-workbook-input.json")
    parser.add_argument("--gate-request", required=True, help="current PAYROLL_EXPORT workflow_gate request JSON")
    parser.add_argument("--calculation-response", required=True, help="原样保存的 get_payroll_calculation DTO JSON")
    parser.add_argument("--person-matches", required=True, help="与当前 attendance.json 摘要绑定的 person-matches.json")
    parser.add_argument("--submission-request", required=True, help="本批次原始 BATCH_SNAPSHOT submit_attendance 请求 JSON")
    parser.add_argument("--output", required=True, help="new .xlsx output path; must not already exist")
    args = parser.parse_args(argv)
    try:
        result = export(Path(args.input), Path(args.gate_request), Path(args.output),
                        Path(args.calculation_response), Path(args.person_matches), Path(args.submission_request))
    except ExportError as exc:
        result = {"ok": False, "status": "EXPORT_BLOCKED", "error": str(exc) or "OUTPUT_WRITE_FAILED"}
    except OSError:
        result = {"ok": False, "status": "EXPORT_BLOCKED", "error": "OUTPUT_WRITE_FAILED"}
    except (KeyError, TypeError, IndexError, AttributeError, OverflowError, ValueError):
        result = {"ok": False, "status": "EXPORT_BLOCKED", "error": "INPUT_CONTRACT_INVALID"}
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0 if result.get("ok") is True else 2


if __name__ == "__main__":
    sys.exit(main())
